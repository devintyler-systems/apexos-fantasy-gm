"""Markdown report -> Excel workbook: one sheet per '## ' section, markdown tables become real tables. Usage: python report_xlsx.py docs/weekly/SLATE_FULL_YYYY_WNN.md out.xlsx"""
import re,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
src,out=sys.argv[1:3]
lines=open(src,encoding="utf-8").read().split("\n")
secs=[];cur=None;pre=[]
for l in lines:
    if l.startswith("## "):
        cur=dict(title=l[3:].strip(),items=[]);secs.append(cur)
    elif cur is None: pre.append(l)
    else: cur["items"].append(l)
wb=Workbook();ws=wb.active;ws.title="Read me"
ws.column_dimensions["A"].width=140
for i,l in enumerate([x for x in pre if x.strip()],1):
    ws.cell(i,1,re.sub(r"\*\*|`","",l.lstrip("# "))).alignment=Alignment(wrap_text=True,vertical="top")
ws["A1"].font=Font(bold=True,size=14)
hf=PatternFill("solid",fgColor="1F3864")
def clean(s): return re.sub(r"\*\*|`","",s.strip())
def num(v):
    m=re.fullmatch(r"([+-]?\d+(?:\.\d+)?)(%?)",v)
    if not m: return v
    x=float(m.group(1)); return x/100 if m.group(2) else (int(x) if x==int(x) and "." not in m.group(1) else x)
used=set()
for s in secs:
    name=re.sub(r"[\[\]:*?/\\]","",re.sub(r"^\d+b?\.\s*","",s["title"]))[:28]
    n=re.match(r"^(\d+b?)\.",s["title"]); name=(n.group(1)+" "+name)[:31] if n else name[:31]
    ws=wb.create_sheet(name); r=1
    ws.cell(r,1,clean(s["title"])).font=Font(bold=True,size=13); r+=2
    items=s["items"];i=0
    while i<len(items):
        l=items[i]
        if l.startswith("|") and i+1<len(items) and re.match(r"^\|[\s\-|:]+\|$",items[i+1]):
            hdr=[clean(c) for c in l.strip().strip("|").split("|")]
            for c,h in enumerate(hdr,1):
                x=ws.cell(r,c,h);x.font=Font(bold=True,color="FFFFFF");x.fill=hf
            r+=1;i+=2
            while i<len(items) and items[i].startswith("|"):
                cells=[clean(c) for c in items[i].strip().strip("|").split("|")]
                for c,v in enumerate(cells,1):
                    x=ws.cell(r,c,num(v))
                    if isinstance(x.value,float) and v.endswith("%"): x.number_format="0.0%"
                r+=1;i+=1
            r+=1
        else:
            if l.strip():
                t=clean(l.lstrip("# "))
                x=ws.cell(r,1,t); x.alignment=Alignment(wrap_text=False)
                if l.startswith("###") or (l.startswith("**") and l.endswith("**")): x.font=Font(bold=True)
                r+=1
            i+=1
    for c in range(1,ws.max_column+1):
        w=max((len(str(ws.cell(rr,c).value or "")) for rr in range(1,ws.max_row+1) if ws.cell(rr,1).value is None or len(str(ws.cell(rr,1).value))<120),default=8)
        ws.column_dimensions[get_column_letter(c)].width=min(max(10,w+2),60)
    ws.freeze_panes=None
wb.save(out)
