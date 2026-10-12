from paths import OUT, NFLV, DELIV, CONF
import os
import pandas as pd, numpy as np
from openpyxl.styles import Font, PatternFill, Alignment
P=pd.read_pickle(OUT+"proj_all.pkl")
O=DELIV
SLATE=os.environ.get("APEX_SLATE_DATE",""); SUF=f"_{SLATE}" if SLATE else ""
_E=pd.read_csv(OUT+"w4_game_environment.csv"); _D={f"{a}@{h}":d for a,h,d in zip(_E.away,_E.home,_E.date)}
P["gdate"]=P.game.map(_D)
P=P[~(P.played | ((P.gdate!=SLATE) if SLATE else False))].copy()      # played games and games outside the requested slate date never appear in a deliverable
WK=os.environ.get("APEX_TARGET_WEEK","")
TW=os.environ.get("APEX_THROUGH_WEEK","") or (int(WK)-1 if WK else "")
TRACK=os.environ.get("APEX_ENV","model")
TRACK_LABEL="Market-anchored (nflverse spread/total as the team-points prior)" if TRACK=="market" else "Model-only (ratings + weather, no lines)"
LINES_NOTE="Spread and total from the nflverse schedule feed the environment only; no player prop odds are used." if TRACK=="market" else "No sportsbook lines or odds are used anywhere in this workbook."
from load import norm_name
_SEASON=int(os.environ.get("APEX_SEASON",2026)); _TW=int(os.environ.get("APEX_THROUGH_WEEK",99))
_st=pd.read_parquet(NFLV+f"stats_player_week_{_SEASON}.parquet"); _st=_st[(_st.week<=_TW)&(_st.season_type=="REG")].copy(); _st["k"]=_st.player_display_name.map(norm_name)
_car=_st[_st.carries>0].groupby(["k","team"]).carries.agg(["mean","size"]); _tg=_st[_st.targets>0].groupby(["k","team"]).targets.agg(["mean","size"])
def uflag(x):
    """Auto flag: projection well above the player's own workload in his games so far (candidates for upper-range reading, never a pick)."""
    k=(norm_name(x.player),x.team); out=[]
    if k in _car.index and x.rush_att_mean>=3 and _car.loc[k,"size"]>=2 and x.rush_att_mean>=1.4*_car.loc[k,"mean"]: out.append(f"projected {x.rush_att_mean:.1f} carries vs {_car.loc[k,'mean']:.1f} per game so far")
    if k in _tg.index and x.tgt_mean>=2 and _tg.loc[k,"size"]>=2 and x.tgt_mean>=1.4*_tg.loc[k,"mean"]: out.append(f"projected {x.tgt_mean:.1f} targets vs {_tg.loc[k,'mean']:.1f} per game so far")
    return ("UPPER-RANGE: "+"; ".join(out)) if out else ""
pm=lambda x:f"{x:.0f}"
def rng(a,b,d=0): return f"{a:.{d}f} to {b:.{d}f}"
def flag(r):
    parts=[x for x in (str(r.status) if isinstance(r.status,str) and r.status else "", uflag(r) if r.pos in ("RB","WR","TE") else "") if x]
    return "; ".join(parts)
# ---------- QB ----------
q=P[(P.pos=="QB")&(~P.proxy)&(P.p_active>=0.3)].copy().sort_values("pass_yds_mean",ascending=False)
QB=pd.DataFrame({"Player":q.player,"Team":q.team,"Opp":q.opp,"Availability %":(q.p_active*100).round(0),
 "Pass Cmp":q.pass_cmp_mean.round(1),"Cmp range (P25-P75)":[rng(a,b) for a,b in zip(q.pass_cmp_p25,q.pass_cmp_p75)],"Cmp conf":q.pass_cmp_conf,
 "Pass Yds":q.pass_yds_mean.round(0),"Yds range (P25-P75)":[rng(a,b) for a,b in zip(q.pass_yds_p25,q.pass_yds_p75)],"Yds conf":q.pass_yds_conf,
 "Pass TD":q.pass_td_mean.round(2),"TD range (P25-P75)":[rng(a,b) for a,b in zip(q.pass_td_p25,q.pass_td_p75)],"TD conf":q.pass_td_conf,"Notes":[("Questionable: starts in %.0f%% of simulated games, backup in the rest. " % (100*r.p_active) if r.p_active<0.97 else "")+flag(r) for r in q.itertuples()]})
# ---------- RB ----------
r=P[(P.pos=="RB")&((P.rush_att_mean>=3)|(P.tgt_mean>=2))].copy(); r["scrim"]=r.rush_yds_mean+r.rec_yds_mean; r=r.sort_values("scrim",ascending=False)
RB=pd.DataFrame({"Player":r.player,"Team":r.team,"Opp":r.opp,"Availability %":(r.p_active*100).round(0),
 "Rush Yds":r.rush_yds_mean.round(0),"Rush range (P25-P75)":[rng(a,b) for a,b in zip(r.rush_yds_p25,r.rush_yds_p75)],"Rush conf":r.rush_yds_conf,
 "Rec Yds":r.rec_yds_mean.round(0),"Rec range (P25-P75)":[rng(a,b) for a,b in zip(r.rec_yds_p25,r.rec_yds_p75)],"Rec conf":r.rec_yds_conf,
 "Rush TD (exp)":r.rush_td_mean.round(2),"Rec TD (exp)":r.rec_td_mean.round(2),"Total TD (exp)":r.exp_td.round(2),"Anytime TD % (if active)":(r.any_td_pct*100).round(1),"TD conf":r.td_conf,"Notes":[flag(x) for x in r.itertuples()]})
# ---------- WR/TE ----------
w=P[(P.pos.isin(["WR","TE"]))&(P.tgt_mean>=2)].copy().sort_values("rec_yds_mean",ascending=False)
WT=pd.DataFrame({"Player":w.player,"Pos":w.pos,"Team":w.team,"Opp":w.opp,"Availability %":(w.p_active*100).round(0),
 "Rec Yds":w.rec_yds_mean.round(0),"Yds range (P25-P75)":[rng(a,b) for a,b in zip(w.rec_yds_p25,w.rec_yds_p75)],"Yds conf":w.rec_yds_conf,
 "Receptions":w.rec_mean.round(1),"Rec range (P25-P75)":[rng(a,b) for a,b in zip(w.rec_p25,w.rec_p75)],"Rec conf":w.rec_conf,
 "TD (exp)":w.exp_td.round(2),"Anytime TD % (if active)":(w.any_td_pct*100).round(1),"TD conf":w.td_conf,"Notes":[flag(x) for x in w.itertuples()]})
# ---------- Anytime TD ----------
A=P.copy(); A["role"]=np.where(A.pos=="QB","QB rushing",np.where(A.pos=="RB",("RB: projected "+A.rush_att_mean.fillna(0).round(1).astype(str)+" carries, "+A.tgt_mean.fillna(0).round(1).astype(str)+" targets"),A.pos+": projected "+A.tgt_mean.fillna(0).round(1).astype(str)+" targets"))
A["avail_adj_pct"]=A.any_td_pct_wt*100
def tdrow(x): return {"Player":x.player,"Pos":x.pos,"Team":x.team,"Opp":x.opp,"Game":x.game,"Anytime TD % (availability-adj)":round(x.any_td_pct_wt*100,1),"If-active %":round(x.any_td_pct*100,1),"Exp TDs":round(x.exp_td,2),"Reliability index (1-100, not a probability)":x.td_conf,"Role (projected usage)":x.role,"Notes":flag(x)}
games=list(dict.fromkeys(P.game))
tops=[]
for gm in games:
    x=A[A.game==gm].sort_values("any_td_pct_wt",ascending=False).head(5)
    for i,xx in enumerate(x.itertuples(),1): tops.append(dict(Game=gm.replace("@"," @ "),Rank=i,**{k:v for k,v in tdrow(xx).items() if k!="Game"}))
T5=pd.DataFrame(tops)
T20=A[~A.played].sort_values("any_td_pct_wt",ascending=False).head(20)
T20=pd.DataFrame([dict(Rank=i,**tdrow(x)) for i,x in enumerate(T20.itertuples(),1)])
METHOD=pd.DataFrame({"Item":["What this is","Projection","Range","Availability %","Anytime TD %","Confidence (1-100)","  ingredients","  not","Scoring","Skipped / caveats"],
 "Detail":[f"{TRACK_LABEL} Week {WK} projections from 20,000 correlated game simulations per game. {LINES_NOTE}",
 "Mean of the simulated stat, conditional on the player playing.",
 "P25 to P75 of the simulated stat (the middle half of outcomes).",
 "Chance the player is active, from the NFL injury report and practice status via the fitted table in config/weekly/availability_by_status*.csv (by report status x practice x position, 2023-26): Out 0; Doubtful about 1-2%; Questionable about 50% (DNP), 65-67% (Limited) or 75% (Full participation); no game designation, WR/TE about 81% (DNP) / 95% (Full or Limited), RB about 60% (DNP) / 95-99%, QB set to 98% when he is the depth-chart starter. Roster-status Out (IR, released, retired, left team) is 0.",
 "Probability of at least one rushing or receiving TD (passing TDs do not count). Table shows if-active; the TD lists rank by availability-adjusted % (if-active x availability).",
 "Reliability of the projection, not the chance it hits. Score = Availability x (60% tightness + 40% sample/role stability). Tightness = 1 - (P75-P25) / (2 x projection): the narrower the middle-50% range relative to the number, the higher.",
 f"TD conf (all TD columns and lists) is different: it measures the evidence behind the TD estimate = Availability x (50% role/sample stability + 50% red-zone touches observed, full credit at 8 inside-the-20 targets+carries; capped at 85 because TDs are inherently noisy). Passing-TD conf uses the tightness formula. Sample/role stability = Weeks 1-{TW} volume relative to a full {TW}-game role, scaled by games played, and (for pass catchers) by QB availability.",
 "Typical ranges: completions and passing yards highest, then rushing yards and receptions, then receiving yards; TD projections are the lowest by nature (small counts).",
 "Counting stats only; no fantasy-point conversion in these tables.",
 "Games already played and games outside the slate date are removed from every sheet. Availability reflects the NFL injury report as of the data fetch recorded in the run manifest and will move until the official inactives post about 90 minutes before each kickoff. Backup QBs fill in by depth-chart rank then attempts (DECISIONS #9), weighted by the starter's availability. KNOWN LIMITS: (1) official inactives are not applied until a rerun; (2) quarterback identity does not carry to receivers, so a backup projects at or near the starter he replaces; (3) a team with injured RBs can have carries that are assigned to nobody (PHI 2026 W5: about 11 of 25.6); (4) the target-share cap is disabled; (5) backup running backs and receivers who inherit vacated work are projected above their own history and are marked UPPER-RANGE in Notes; (6) every Reliability index is capped at 85 for touchdowns and is not a probability; (7) the played and off-slate games are not in this workbook."]})
with pd.ExcelWriter(O+f"Week{WK}_{TRACK}_Projections{SUF}.xlsx",engine="openpyxl") as xw:
    for name,df in (("Method",METHOD),("QB",QB),("RB",RB),("WR_TE",WT),("TD Top5 by Game",T5),("TD Top20",T20)):
        df=df.rename(columns=lambda c: c.replace(" conf"," reliability index (not a probability)") if str(c).endswith(" conf") else c)
        df.to_excel(xw,sheet_name=name,index=False); ws=xw.sheets[name]
        for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F3A5F"); c.alignment=Alignment(wrap_text=True,vertical="center")
        for col in ws.columns:
            L=max(len(str(c.value)) if c.value is not None else 0 for c in col); ws.column_dimensions[col[0].column_letter].width=min(max(10,L+2),70 if name=="Method" else 34)
        ws.freeze_panes="B2" if name not in ("Method",) else "A2"
        if name=="Method":
            for row in ws.iter_rows(min_row=2):
                for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
for n,df in (("QB",QB),("RB",RB),("WR_TE",WT),("TD_Top5_by_game",T5),("TD_Top20",T20)): df.rename(columns=lambda c: c.replace(" conf"," reliability index (not a probability)") if str(c).endswith(" conf") else c).to_csv(O+f"model_proj_{n}{SUF}.csv",index=False)
print(len(QB),len(RB),len(WT),len(T5),len(T20))
pd.to_pickle(dict(QB=QB,RB=RB,WT=WT,T5=T5,T20=T20),OUT+"proj_tables.pkl")
pd.set_option("display.width",250)
print(T20[["Rank","Player","Team","Opp","Anytime TD % (availability-adj)","If-active %","Exp TDs","Notes"]].to_string(index=False))
